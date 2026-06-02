from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

doc = Document()

# Tiêu đề
title = doc.add_heading('BÁO CÁO BÀI TẬP LỚN', 0)
title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

doc.add_paragraph('Môn học: Lập trình mạng (Network Programming)').alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
doc.add_paragraph('Đề tài: Hệ thống Đấu giá Trực tuyến (Online Bidding System)').alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

doc.add_heading('1. Giới thiệu mục tiêu và phạm vi thực hiện', level=1)
doc.add_paragraph('Mục tiêu: Xây dựng một ứng dụng Client-Server hỗ trợ người dùng đăng bán, đấu giá sản phẩm theo thời gian thực dựa trên giao thức TCP Socket. Ứng dụng mô phỏng quy trình đấu giá thực tế gồm: mở phiên, trả giá liên tục, và thanh toán tự động khi kết thúc.')
doc.add_paragraph('Phạm vi: Hệ thống có 3 vai trò chính: Admin (Quản trị hệ thống, duyệt sản phẩm/tiền), Seller (Người đăng bán), và Bidder (Người mua). Hệ thống chạy trên môi trường local hoặc mạng LAN/WAN, hỗ trợ đa nền tảng nhờ sử dụng Java.')

doc.add_heading('2. Kiến trúc tổng thể của hệ thống', level=1)
doc.add_paragraph('Hệ thống sử dụng mô hình Client-Server với giao thức TCP/IP Socket:')
doc.add_paragraph('- Server: Chịu trách nhiệm nhận kết nối từ nhiều Client (Sử dụng Multi-threading ThreadPool). Xử lý logic nghiệp vụ, giao tiếp với cơ sở dữ liệu MySQL (qua JDBC), và phát sóng (Broadcast) dữ liệu realtime đến các Client. Server cũng tích hợp một Scheduler để lập lịch mở/đóng các phiên đấu giá một cách tự động.')
doc.add_paragraph('- Client: Là ứng dụng Desktop phát triển bằng JavaFX. Giao tiếp với Server thông qua việc gửi các đối tượng Request/Response đã được Serialize. Client có các luồng riêng biệt để lắng nghe (Listen) dữ liệu realtime đẩy từ Server (ví dụ: thông báo, giá mới).')
doc.add_paragraph('- Database: Cơ sở dữ liệu quan hệ MySQL lưu trữ thông tin User, Sản phẩm, Phiên đấu giá, Lịch sử Bid, Giao dịch ví tiền và Thông báo.')

doc.add_heading('3. Chức năng đạt được theo barem điểm', level=1)

func1 = doc.add_paragraph()
func1.add_run('1. Xác thực và Phân quyền: ').bold = True
func1.add_run('Hệ thống hỗ trợ Đăng ký, Đăng nhập. Có 3 vai trò: BIDDER, SELLER, ADMIN. Admin có Dashboard riêng. Hướng giải quyết: Dùng bảng `users` kết hợp bảng `user_profiles`, phân biệt quyền qua cột role.')

func2 = doc.add_paragraph()
func2.add_run('2. Quản lý Sản phẩm & Phiên đấu giá: ').bold = True
func2.add_run('Người bán thêm sản phẩm. Admin duyệt. Sau khi duyệt, phiên chuyển sang ACTIVE và tự động chuyển sang RUNNING khi đến giờ. Hướng giải quyết: Sử dụng Timer/Scheduler trên Server để quét các phiên định kỳ, đảm bảo tính tự động hoàn toàn mà không cần can thiệp thủ công.')

func3 = doc.add_paragraph()
func3.add_run('3. Đấu giá thời gian thực (Real-time Bidding): ').bold = True
func3.add_run('Nhiều người dùng có thể đặt giá cùng lúc. Giá mới nhất tự động cập nhật lên màn hình tất cả người đang xem. Hướng giải quyết: Dùng TCP Socket và kỹ thuật Broadcast từ Server tới danh sách các ClientHandler đang quan tâm phiên đấu giá đó. Từ khóa `synchronized` được dùng để tránh Race-condition khi ghi dữ liệu ra luồng.')

func4 = doc.add_paragraph()
func4.add_run('4. Ví điện tử & Thanh toán tự động: ').bold = True
func4.add_run('Người dùng tạo yêu cầu nạp/rút tiền (Admin duyệt). Khi phiên kết thúc, tiền tự động trừ ở người thắng và cộng cho người bán. Hướng giải quyết: Đóng gói giao dịch trong Transaction của MySQL (setAutoCommit(false)) để đảm bảo tính toàn vẹn dữ liệu dòng tiền (ACID).')

func5 = doc.add_paragraph()
func5.add_run('5. Thông báo (Notification): ').bold = True
func5.add_run('Hệ thống chủ động đẩy thông báo cho người dùng khi có sự kiện (Trúng đấu giá, Bị từ chối sản phẩm, Tiền được duyệt). Hướng giải quyết: Server chủ động send packet NOTIFICATION cho Client đang online. Client offline sẽ được lưu vào DB và nhận khi đăng nhập.')

doc.add_heading('4. Phân chia công việc', level=1)
table = doc.add_table(rows=1, cols=3)
table.style = 'Table Grid'
hdr_cells = table.rows[0].cells
hdr_cells[0].text = 'Họ và Tên'
hdr_cells[1].text = 'Nhiệm vụ chính'
hdr_cells[2].text = 'Mức độ đóng góp'

row_cells = table.add_row().cells
row_cells[0].text = 'Thành viên 1'
row_cells[1].text = 'Thiết kế Database, Code Server Socket, Xử lý Đa luồng & Scheduler'
row_cells[2].text = '100%'

row_cells = table.add_row().cells
row_cells[0].text = 'Thành viên 2'
row_cells[1].text = 'Thiết kế UI/UX JavaFX, Client Socket, Ghép nối API'
row_cells[2].text = '100%'

doc.add_paragraph('\n(Note: Bạn có thể chỉnh sửa lại tên thành viên và % đóng góp cho phù hợp thực tế)')

doc.save('Bao_cao_BTL.docx')
print("Đã tạo file Bao_cao_BTL.docx thành công!")
